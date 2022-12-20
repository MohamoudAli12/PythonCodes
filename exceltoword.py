import openpyxl
import docx

workbook= openpyxl.load_workbook('malik.xlsx')
sheet =workbook.get_sheet_by_name('Sheet2')
#cell =sheet['C3']
for i in range (2,1594):

    doc=docx.Document('word2.docx')
    doc.add_paragraph(sheet.cell(row=i,column=4).value)
    doc.save("word2.docx")
print("DONE!!!")
